"""把说明书渲染成 Word 文件（需要 python-docx）。"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from gongke_core.report import Doc


def _font(run, size=None, bold=None, name="宋体"):
    run.font.name = name
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def write(doc: Doc, path: str | Path) -> None:
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    style = d.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    title_style = d.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    for border in title_style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)

    p = d.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(doc.title), 18, True, "黑体")

    for b in doc.blocks:
        if b.kind in ("h1", "h2"):
            p = d.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(12 if b.kind == "h1" else 6)
            _font(p.add_run(b.text), 15 if b.kind == "h1" else 13, True, "黑体")
        elif b.kind == "p":
            p = d.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            _font(p.add_run(b.text))
        elif b.kind == "eq":
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            _font(p.add_run(b.text), name="宋体")
        elif b.kind == "note":
            p = d.add_paragraph()
            _font(p.add_run("注：" + b.text), 10.5)
        elif b.kind == "table":
            t = d.add_table(rows=1, cols=len(b.headers))
            t.style = "Table Grid"
            t.autofit = False
            # 十二列疲劳校核表为截面名称留出空间，数值列使用紧凑字号。
            widths = ([2.8, 1, 1.2, 1.2, 1.5, 1, 1, 1.2, 1.2, 1.3, 1.3, 1.3]
                      if len(b.headers) == 12 else [16 / len(b.headers)] * len(b.headers))
            if b.headers == ["项目", "结论", "说明"]:
                widths = [5.5, 1.5, 9]
            elif len(b.headers) == 7 and b.headers[1] == "轴径 / mm":
                widths = [3.8, 1.5, 3.5, 1, 2.3, 2.3, 1.6]
            elif len(b.headers) == 9 and b.headers[1] == "代号":
                widths = [3, 1.3, 1.6, 1.4, 1.4, 1.5, 1.4, 2.5, 1.9]
            size = 8.5 if len(b.headers) == 12 else 10.5
            for col, width in zip(t.columns, widths):
                col.width = Cm(width)
            header = OxmlElement("w:tblHeader")
            t.rows[0]._tr.get_or_add_trPr().append(header)
            for c, h in zip(t.rows[0].cells, b.headers):
                c.text = ""
                _font(c.paragraphs[0].add_run(h.replace("/", "\n") if len(b.headers) == 12 else h), size, True)
            for row in b.rows:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = ""
                    # 成对系数分行，避免 Word 把小数末位挤到下一行。
                    value = v.replace("/", "\n") if len(b.headers) == 12 else v
                    _font(c.paragraphs[0].add_run(value), size)
            for row in t.rows:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
                    if len(b.headers) == 12:
                        margins = OxmlElement("w:tcMar")
                        for side in ("left", "right"):
                            margin = OxmlElement(f"w:{side}")
                            margin.set(qn("w:w"), "45")
                            margin.set(qn("w:type"), "dxa")
                            margins.append(margin)
                        cell._tc.get_or_add_tcPr().append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
            for cell in t.rows[0].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        elif b.kind == "img":
            d.add_picture(io.BytesIO(b.image), width=Cm(15.5))
            d.paragraphs[-1].paragraph_format.keep_with_next = True
            d.paragraphs[-1].paragraph_format.space_before = Pt(8)
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(b.text), 10.5)
    d.save(str(path))
